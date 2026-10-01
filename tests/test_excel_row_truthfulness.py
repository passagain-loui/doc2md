"""Milestone 7 item 3 - Excel/CSV row-count truthfulness.

`rows_detected` must be the REAL row count, not a value frozen at
`max_rows + 1` the moment truncation was first detected - a 100,000-row
file capped at 100 must report 100,000 detected, not 101. `rows_detected_is_exact`
records whether the count is real (True) or a lower bound (never emitted by
these two engines any more, but the field and its GUI/JSON wording must
still degrade correctly if some future engine sets it False).
"""

from __future__ import annotations

from doc2md.engine.excel_engine import ExcelEngine
from doc2md.gui.main_window import _format_rows_detected
from doc2md.core.quality import QualityMetrics


LARGE_N = 50_000  # substantially larger than any max_rows used below


def _write_large_csv(path, n_data_rows=LARGE_N):
    lines = ["id,label"] + [f"{i},row{i}" for i in range(n_data_rows)]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _write_large_xlsx(path, n_data_rows=LARGE_N, sheet_name="Data"):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["id", "label"])
    for i in range(n_data_rows):
        ws.append([i, f"row{i}"])
    wb.save(str(path))
    return path


def test_csv_reports_exact_row_count_far_beyond_max_rows(tmp_path):
    p = _write_large_csv(tmp_path / "big.csv")
    output = ExcelEngine().convert_structured(p, {"max_rows": 100, "sample_rows": 5})

    assert output.metrics.truncated is True
    assert output.metrics.rows_detected == LARGE_N + 1  # header + data rows
    assert output.metrics.rows_detected_is_exact is True
    assert output.metrics.rows_exported == 5
    # the markdown note must not claim a lower bound now that it's exact
    assert ">=" not in output.markdown
    assert f"{LARGE_N + 1:,} rows detected" in output.markdown


def test_xlsx_reports_exact_row_count_far_beyond_max_rows(tmp_path):
    p = _write_large_xlsx(tmp_path / "big.xlsx")
    output = ExcelEngine().convert_structured(p, {"max_rows": 100, "sample_rows": 5})

    assert output.metrics.truncated is True
    assert output.metrics.rows_detected == LARGE_N + 1
    assert output.metrics.rows_detected_is_exact is True
    assert output.metrics.rows_exported == 5


def test_xlsx_multi_sheet_exact_counts_mixed_truncation(tmp_path):
    import openpyxl

    p = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    small = wb.active
    small.title = "Small"
    small.append(["a"])
    for i in range(5):
        small.append([i])  # 6 total rows, under max_rows

    big = wb.create_sheet("Big")
    big.append(["b"])
    for i in range(LARGE_N):
        big.append([i])  # LARGE_N + 1 total rows, over max_rows
    wb.save(str(p))

    output = ExcelEngine().convert_structured(p, {"max_rows": 100, "sample_rows": 5})

    assert output.metrics.sheets_detected == 2
    assert output.metrics.rows_detected == 6 + (LARGE_N + 1)
    assert output.metrics.rows_detected_is_exact is True
    assert output.metrics.truncated is True  # any sheet truncated -> True overall


def test_untruncated_file_is_also_marked_exact(tmp_path):
    p = _write_large_csv(tmp_path / "small.csv", n_data_rows=10)
    output = ExcelEngine().convert_structured(p, {})

    assert output.metrics.truncated is False
    assert output.metrics.rows_detected == 11
    assert output.metrics.rows_detected_is_exact is True


# --- JSON report shape -------------------------------------------------------------


def test_quality_report_carries_the_exact_flag(tmp_path):
    from doc2md.core.converter import Converter
    from doc2md.core.quality import build_report

    p = _write_large_csv(tmp_path / "big.csv")
    result = Converter(options={"max_rows": 100, "sample_rows": 5}).convert_file(p)
    report = build_report(result)

    assert report["rows_detected"] == LARGE_N + 1
    assert report["rows_detected_is_exact"] is True


# --- GUI wording (pure function, no Qt needed) --------------------------------------


def test_gui_wording_shows_plain_count_when_exact():
    quality = QualityMetrics(rows_detected=100_001, rows_detected_is_exact=True)
    assert _format_rows_detected(quality) == "100,001"


def test_gui_wording_shows_lower_bound_marker_when_not_exact():
    quality = QualityMetrics(rows_detected=101, rows_detected_is_exact=False)
    assert _format_rows_detected(quality) == "101+"


def test_gui_wording_omits_marker_when_exactness_unknown():
    """A legacy, not-yet-migrated engine that sets rows_detected without
    ever populating rows_detected_is_exact must not be shown as a false
    lower bound OR a false exact count - it just shows the number plainly."""
    quality = QualityMetrics(rows_detected=42, rows_detected_is_exact=None)
    assert _format_rows_detected(quality) == "42"


def test_gui_wording_none_when_nothing_detected():
    assert _format_rows_detected(QualityMetrics()) is None
