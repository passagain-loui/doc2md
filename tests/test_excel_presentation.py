"""How spreadsheet content is presented: titles, dates, percentages, hidden sheets."""

from __future__ import annotations

import datetime as dt

import openpyxl
import pytest
from typer.testing import CliRunner

from doc2md.cli.main import app
from doc2md.engine.excel_engine import ExcelEngine, format_cell_value

runner = CliRunner()


def _convert(path, **options):
    return ExcelEngine().convert(path, options)


def test_a_title_row_is_text_not_the_table_header(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "Internal Audit Result Report "
    sheet.append(["No.", "Ref", "Department", "Result"])
    sheet.append([1, "CAR26/01", "Customer Service ", "Minor"])
    path = tmp_path / "t.xlsx"
    workbook.save(path)

    output = _convert(path)

    assert "Internal Audit Result Report\n" in output
    assert "| No. | Ref | Department | Result |" in output
    assert "col2" not in output
    assert "| 1 | CAR26/01 | Customer Service | Minor |" in output


def test_a_sparse_two_value_row_is_left_as_the_header(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Value"])
    sheet.append(["a", 1])
    path = tmp_path / "t.xlsx"
    workbook.save(path)

    assert "| Name | Value |" in _convert(path)


def test_midnight_datetimes_are_plain_dates(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["a", "b", "c"])
    sheet.append([dt.datetime(2026, 8, 17), dt.datetime(2026, 8, 17, 13, 30, 5), 1])
    path = tmp_path / "t.xlsx"
    workbook.save(path)

    output = _convert(path)

    assert "2026-08-17 |" in output
    assert "2026-08-17 13:30:05" in output
    assert "00:00:00" not in output


@pytest.mark.parametrize(
    ("value", "number_format", "expected"),
    [
        (-0.1724137931034483, "0%", "-17%"),
        (0.256, "0.0%", "25.6%"),
        (3.0, "General", "3"),
        (0.1 + 0.2, "General", "0.3"),
        (12.5, "0.00", "12.50"),
        ("  padded  ", "General", "padded"),
        ("   ", "General", None),
        (7, "General", 7),
    ],
)
def test_cell_value_formatting(value, number_format, expected):
    assert format_cell_value(value, number_format) == expected


def _workbook_with_hidden_sheet(path):
    workbook = openpyxl.Workbook()
    shown = workbook.active
    shown.title = "Shown"
    shown.append(["a", "b", "c"])
    shown.append([1, 2, 3])
    secret = workbook.create_sheet("Old Archive ")
    secret.append(["x", "y", "z"])
    secret.append(["old", "data", "here"])
    secret.sheet_state = "hidden"
    workbook.save(path)


def test_hidden_sheets_are_skipped_with_a_note_naming_them(tmp_path):
    path = tmp_path / "t.xlsx"
    _workbook_with_hidden_sheet(path)

    output = _convert(path)

    assert "## Sheet: Shown" in output
    assert "old" not in output
    assert "1 hidden sheet(s) were not converted: `Old Archive`" in output


def test_include_hidden_sheets_converts_everything(tmp_path):
    path = tmp_path / "t.xlsx"
    _workbook_with_hidden_sheet(path)

    output = _convert(path, include_hidden_sheets=True)

    assert "## Sheet: Old Archive" in output
    assert "hidden sheet(s) were not converted" not in output


def test_cli_include_hidden_flag(tmp_path):
    path = tmp_path / "t.xlsx"
    _workbook_with_hidden_sheet(path)

    plain = runner.invoke(app, ["convert", str(path), "--stdout"])
    full = runner.invoke(app, ["convert", str(path), "--stdout", "--include-hidden"])

    assert "Old Archive" in plain.output  # named in the note only
    assert "| old | data | here |" not in plain.output
    assert "| old | data | here |" in full.output


def _merged_header_workbook(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Plan"
    sheet["A1"] = "Quality plan"
    sheet.merge_cells("A1:F1")
    sheet["A2"], sheet["B2"], sheet["D2"], sheet["F2"] = "No.", "Plan", "Actual", "Group"
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:C2")
    sheet.merge_cells("D2:E2")
    sheet.merge_cells("F2:F3")
    for column, label in zip("BCDE", ("Q1", "Q2", "Q1", "Q2")):
        sheet[f"{column}3"] = label
    sheet.append([1, 10, 11, 12, 13, "Group A"])
    sheet.append([2, 20, 21, 22, 23, None])
    sheet.append([3, 30, 31, 32, 33, None])
    sheet.merge_cells("F4:F6")
    workbook.save(path)


def test_multi_row_merged_headers_become_one_labelled_header(tmp_path):
    path = tmp_path / "t.xlsx"
    _merged_header_workbook(path)

    output = _convert(path)

    assert "Quality plan\n" in output
    assert "| No. | Plan / Q1 | Plan / Q2 | Actual / Q1 | Actual / Q2 | Group |" in output
    assert "col2" not in output


def test_a_vertical_merge_repeats_its_value_on_every_row_it_covers(tmp_path):
    path = tmp_path / "t.xlsx"
    _merged_header_workbook(path)

    output = _convert(path)

    assert "| 1 | 10 | 11 | 12 | 13 | Group A |" in output
    assert "| 2 | 20 | 21 | 22 | 23 | Group A |" in output
    assert "| 3 | 30 | 31 | 32 | 33 | Group A |" in output


def test_a_sheet_without_merges_is_unchanged(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["a", "b", "c"])
    sheet.append([1, 2, 3])
    path = tmp_path / "t.xlsx"
    workbook.save(path)

    assert "| a | b | c |" in _convert(path)


def test_unreadable_merge_data_falls_back_to_plain_conversion(tmp_path, monkeypatch):
    path = tmp_path / "t.xlsx"
    _merged_header_workbook(path)
    monkeypatch.setattr(
        "doc2md.engine.excel_engine.sheet_xml_paths",
        lambda source: (_ for _ in ()).throw(KeyError("broken")),
    )

    output = _convert(path)

    assert "Group A" in output
