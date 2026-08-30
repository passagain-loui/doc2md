"""DOCX engine: order-aware block iteration with structure and inline styling.

What is preserved, and why each one needs explicit handling:

* **Headings** - ``Heading 1..9`` map to ``#``..``######``. Word's localized
  style names (Thai templates ship ``หัวเรื่อง 1``) are matched by the
  underlying ``style_id`` as well as the display name.
* **Lists** - bullet vs numbered is read from the style, and the indent level
  from ``w:numPr/w:ilvl`` so nested lists keep their depth.
* **Bold / italic** - applied per run, not per paragraph, so a sentence with
  one bold phrase does not become entirely bold. Runs are merged before
  emphasis is applied, otherwise Word's habit of splitting a word into several
  runs produces ``**ต**​**ัวหนา**`` and renders as literal asterisks.
* **Tables** - handed to :mod:`doc2md.core.tables`, which pads ragged rows and
  escapes pipes so Thai tables keep their columns.
"""

from __future__ import annotations

import re
from pathlib import Path

from doc2md.core.errors import ConversionError, EngineUnavailableError
from doc2md.core.router import FileKind
from doc2md.core.tables import render_table
from doc2md.engine.base import BaseEngine

_HEADING_STYLE_RE = re.compile(r"(?:heading|หัวเรื่อง)\s*(\d+)", re.IGNORECASE)
_MAX_LIST_DEPTH = 5


class DocxEngine(BaseEngine):
    name = "docx"
    supported_kinds = (FileKind.DOCX,)
    requires_process_isolation = False

    def convert(self, source: Path, options: dict) -> str:
        self.validate_source(source)
        try:
            import docx
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise EngineUnavailableError(
                "DOCX backend missing: pip install -r requirements.txt (python-docx)"
            ) from exc

        try:
            document = docx.Document(str(source))
        except Exception as exc:
            raise ConversionError(
                f"Corrupted or unreadable DOCX: {source} ({exc})"
            ) from exc

        def iter_blocks():
            for child in document.element.body.iterchildren():
                if child.tag == qn("w:p"):
                    yield Paragraph(child, document)
                elif child.tag == qn("w:tbl"):
                    yield Table(child, document)

        emphasis = bool(options.get("inline_styles", True))
        parts: list[str] = [f"# {source.name}", ""]
        try:
            for block in iter_blocks():
                if isinstance(block, Paragraph):
                    line = self._paragraph_to_markdown(block, emphasis=emphasis)
                    if line:
                        parts.append(line)
                        if line.startswith("#"):
                            parts.append("")
                else:
                    table_md = self._table_to_markdown(block)
                    if table_md:
                        parts.extend(["", table_md, ""])
        except ConversionError:
            raise
        except Exception as exc:
            raise ConversionError(f"DOCX conversion failed: {source} ({exc})") from exc

        return "\n".join(parts).strip() + "\n"

    # ------------------------------------------------------------- paragraphs

    @classmethod
    def _paragraph_to_markdown(cls, paragraph, *, emphasis: bool = True) -> str:
        text = cls._inline_text(paragraph, emphasis=emphasis)
        if not text:
            return ""

        style_name, style_id = cls._style_identity(paragraph)
        probe = f"{style_name} {style_id}"

        heading = _HEADING_STYLE_RE.search(probe)
        if heading:
            level = min(max(int(heading.group(1)), 1), 6)
            return f"{'#' * level} {cls._plain(paragraph)}"
        if "title" in probe or "subtitle" in probe:
            level = 2 if "subtitle" in probe else 1
            return f"{'#' * level} {cls._plain(paragraph)}"

        compact = probe.replace(" ", "")
        numbered = cls._has_numbering(paragraph)
        indent = "  " * cls._list_level(paragraph)
        if "listnumber" in compact:
            return f"{indent}1. {text}"
        if "listbullet" in compact or ("listparagraph" in compact and numbered):
            return f"{indent}- {text}"
        if "quote" in probe:
            return f"> {text}"
        if "code" in probe:
            return f"`{cls._plain(paragraph)}`"
        return text

    @staticmethod
    def _style_identity(paragraph) -> tuple[str, str]:
        try:
            style = paragraph.style
            return (style.name or "").lower(), (style.style_id or "").lower()
        except Exception:
            return "", ""

    @staticmethod
    def _plain(paragraph) -> str:
        return " ".join(paragraph.text.split())

    @staticmethod
    def _has_numbering(paragraph) -> bool:
        try:
            pPr = paragraph._p.pPr
            return pPr is not None and pPr.numPr is not None
        except Exception:
            return False

    @staticmethod
    def _list_level(paragraph) -> int:
        """Read the numbering indent level from the paragraph properties."""
        try:
            from docx.oxml.ns import qn

            pPr = paragraph._p.pPr
            if pPr is None or pPr.numPr is None or pPr.numPr.ilvl is None:
                return 0
            level = int(pPr.numPr.ilvl.get(qn("w:val")) or 0)
        except Exception:
            return 0
        return min(max(level, 0), _MAX_LIST_DEPTH)

    @classmethod
    def _inline_text(cls, paragraph, *, emphasis: bool = True) -> str:
        """Merge runs into text, wrapping bold/italic spans in Markdown markers.

        Adjacent runs sharing the same formatting are merged *before* the
        markers are added. Word routinely splits a single styled word across
        several runs (spell-check state, language tags, rsid churn); emitting
        one ``**...**`` pair per run would produce marker soup that renders as
        literal asterisks rather than bold text.
        """
        plain = " ".join(paragraph.text.split())
        if not emphasis or not plain:
            return plain

        spans: list[tuple[bool, bool, str]] = []
        try:
            for run in paragraph.runs:
                if not run.text:
                    continue
                key = (bool(run.bold), bool(run.italic))
                if spans and (spans[-1][0], spans[-1][1]) == key:
                    spans[-1] = (key[0], key[1], spans[-1][2] + run.text)
                else:
                    spans.append((key[0], key[1], run.text))
        except Exception:
            return plain

        if not spans:
            return plain

        out: list[str] = []
        for bold, italic, chunk in spans:
            # Emphasis markers must hug the text: "** bold **" is not bold in
            # any CommonMark renderer. Whitespace is lifted outside the markers.
            lead = chunk[: len(chunk) - len(chunk.lstrip())]
            trail = chunk[len(chunk.rstrip()) :]
            body = " ".join(chunk.split())
            if not body:
                out.append(" ")
                continue
            if bold and italic:
                body = f"***{body}***"
            elif bold:
                body = f"**{body}**"
            elif italic:
                body = f"*{body}*"
            out.append(f"{' ' if lead else ''}{body}{' ' if trail else ''}")

        merged = "".join(out)
        merged = re.sub(r"[ \t]{2,}", " ", merged).strip()
        return merged or plain

    # ----------------------------------------------------------------- tables

    @staticmethod
    def _table_to_markdown(table) -> str:
        rows: list[list[str]] = []
        for row in table.rows:
            # row.cells resolves horizontal merges by repeating the same cell,
            # which is exactly what keeps the column count stable.
            # Cells are handed over raw: render_table escapes exactly once.
            rows.append([cell.text for cell in row.cells])
        return render_table(rows)
